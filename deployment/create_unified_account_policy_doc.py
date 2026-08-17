from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\ALL_project\Soodal_Life_Project")
OUT = ROOT / "docs" / "output" / "수달라이프_고객공급자_통합계정_및_휴대폰본인인증_정책_v1.0.docx"
LOGO = ROOT / "docs" / "source" / "soodal-life-mark.png"

TEAL = "0F766E"
MINT = "E8F7F3"
DARK = "173B36"
GRAY = "667085"
LIGHT = "F5F7F6"
RED = "B42318"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, start=110, bottom=100, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, name="맑은 고딕", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.style = f"Heading {level}"
    p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_run_font(r, size=17 if level == 1 else 13, bold=True, color=TEAL if level == 1 else DARK)
    return p


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.18
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, size=9.5, bold=True, color=DARK)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r, size=9.5, color=DARK)
    else:
        r = p.add_run(text)
        set_run_font(r, size=9.5, color=DARK)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.55)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(item)
        set_run_font(r, size=9.2, color=DARK)


def add_callout(doc, title, body, fill=MINT, title_color=TEAL):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, fill)
    set_cell_margins(cell, 150, 180, 150, 180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=10, bold=True, color=title_color)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(body)
    set_run_font(r, size=9.2, color=DARK)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    # Repeat the header when a table continues on the next page.
    header_props = table.rows[0]._tr.get_or_add_trPr()
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    header_props.append(repeat_header)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, TEAL)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        set_run_font(r, size=8.7, bold=True, color="FFFFFF")
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(values):
            cell = cells[i]
            if row_index % 2:
                shade(cell, LIGHT)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(value))
            set_run_font(r, size=8.4, color=DARK)
            if widths:
                cell.width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


doc = Document()
section = doc.sections[0]
section.top_margin = Cm(1.7)
section.bottom_margin = Cm(1.6)
section.left_margin = Cm(1.8)
section.right_margin = Cm(1.8)

styles = doc.styles
styles["Normal"].font.name = "맑은 고딕"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
styles["Normal"].font.size = Pt(9.5)

# Masthead
mast = doc.add_table(rows=1, cols=2)
mast.autofit = False
mast.columns[0].width = Cm(4.0)
mast.columns[1].width = Cm(13.0)
left, right = mast.rows[0].cells
for c in (left, right):
    shade(c, TEAL)
    set_cell_margins(c, 180, 220, 180, 220)
if LOGO.exists():
    left.paragraphs[0].add_run().add_picture(str(LOGO), width=Cm(2.0))
else:
    r = left.paragraphs[0].add_run("SOODAL LIFE")
    set_run_font(r, size=11, bold=True, color="FFFFFF")
right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = right.paragraphs[0].add_run("운영·개발 기준서")
set_run_font(r, size=9, bold=True, color="D8F4EE")
p = right.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = p.add_run("고객·공급자 통합계정 및 휴대폰 본인인증 정책")
set_run_font(r, size=18, bold=True, color="FFFFFF")
p = right.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = p.add_run("버전 1.0  |  작성일 2026. 8. 16.  |  상태: 구현 기준 확정")
set_run_font(r, size=8.5, color="D8F4EE")

doc.add_paragraph()
add_callout(doc, "결론", "수달 라이프는 계정을 고객용·공급자용으로 분리하지 않고 하나의 로그인 계정에 CUSTOMER와 PROVIDER 역할을 부여한다. 신규 공급자 가입자는 두 역할을 동시에 갖고, 기존 고객은 동일 계정에 공급자 역할을 추가한다. 양쪽 사이트 이동은 기존 보안 쿠키를 사용하며 URL에 토큰을 노출하지 않는다.")

add_heading(doc, "1. 목적과 적용 범위")
add_body(doc, "본 문서는 고객 사이트(soodallife.kr), 공급자 사이트(partner.soodallife.kr), 관리자 사이트(admin.soodallife.kr), API(api.soodallife.kr)에 적용할 통합계정 운영 기준을 정의한다. 회원가입, 역할 추가, 역할별 화면 전환, 휴대폰 본인인증, 기존 계정 보정, 약관 동의와 보안 통제를 포함한다.")
add_bullets(doc, [
    "적용 대상: 신규 고객, 신규 공급자, 기존 고객의 공급자 전환, 기존 공급자의 고객 활동",
    "계정 원칙: 아이디·비밀번호·휴대전화 본인인증 정보는 사용자 계정에 한 번만 보유",
    "업무 분리: 고객 프로필과 공급자 프로필은 역할별 업무·승인 상태를 별도로 관리",
    "외부연동 원칙: 본인인증 업체 연동 전에는 VERIFIED 성공을 생성하지 않고 NOT_INTEGRATED로 차단",
])

add_heading(doc, "2. 최종 권장 계정 모델")
add_table(doc,
          ["구분", "계정·역할 처리", "프로필·승인 상태", "로그인"],
          [
              ["고객 신규 가입", "CUSTOMER 부여", "고객 프로필 생성", "soodallife.kr"],
              ["공급자 신규 가입", "CUSTOMER + PROVIDER 동시 부여", "고객 프로필 + 공급자 프로필 + 지갑 생성; 공급자는 PENDING/INACTIVE", "partner 또는 customer에서 동일 계정 사용"],
              ["기존 고객의 공급자 등록", "동일 사용자에 PROVIDER 추가", "공급자 프로필·지갑 생성; 고객 데이터 유지", "기존 아이디·비밀번호 유지"],
              ["기존 공급자 계정", "CUSTOMER가 없으면 안전하게 추가", "고객 프로필이 없으면 생성; 기존 공급자 데이터 유지", "기존 아이디·비밀번호 유지"],
          ], [3.1, 5.1, 6.6, 3.8])

add_heading(doc, "3. 화면 및 이동 규칙")
add_table(doc,
          ["현재 상태", "고객 사이트 상단 버튼", "이동 대상"],
          [
              ["CUSTOMER만 보유", "공급자 등록", "공급자 등록 화면"],
              ["PROVIDER 등록/심사 중", "공급자 등록 현황", "공급자 온보딩·심사 현황"],
              ["PROVIDER 승인·활성", "공급자 홈", "partner.soodallife.kr/provider"],
              ["공급자 사이트 + CUSTOMER 보유", "고객 홈", "soodallife.kr/customer"],
          ], [4.2, 5.0, 9.0])
add_bullets(doc, [
    "사이트 간 이동은 현재 인증 쿠키를 재사용하며 자동 로그인을 위해 토큰을 URL·쿼리문자열에 넣지 않는다.",
    "공급자 역할 추가 직후 인증 세션을 갱신하여 새 역할이 즉시 반영되도록 한다.",
    "공급자 심사 거절·보류·비활성 상태여도 고객 역할과 고객 서비스 이용은 정지하지 않는다.",
    "로그인 상태 홈 배너는 ‘원하는 서비스를 선택하고 견적 요청을 시작해 보세요.’로 표시한다. 비로그인 상태에서만 로그인 안내를 표시한다.",
])

add_heading(doc, "4. 휴대폰 본인인증 정책")
add_callout(doc, "필수 Gate", "고객 회원가입과 공급자 신규 가입 모두 휴대폰 본인인증 완료가 필수다. 기존 고객이 공급자 역할을 추가할 때도 계정의 휴대전화 인증 상태가 VERIFIED인지 재확인한다.", fill="FFF4E5", title_color=RED)
add_table(doc,
          ["단계", "검증 항목", "실패 시 처리"],
          [
              ["인증 요청", "휴대전화 번호와 인증 거래 생성", "가입 진행 금지"],
              ["인증 결과 수신", "업체 서명·거래번호·만료·재사용 여부 확인", "위조·만료·재사용 차단"],
              ["가입 제출", "인증된 번호와 가입 입력 번호 일치 확인", "불일치 시 400 오류"],
              ["저장", "PhoneVerificationStatusCode=VERIFIED", "외부연동 전에는 NOT_INTEGRATED 유지"],
          ], [3.0, 8.5, 6.0])
add_bullets(doc, [
    "휴대전화 번호는 화면에서 010-1234-5678 형식으로 표시하되 서버에서는 정규화 후 비교한다.",
    "인증 토큰 원문, 주민등록번호, 통신사 인증 비밀값을 로그·URL·Git에 저장하지 않는다.",
    "현재 외부 본인인증 연동 전 상태에서는 성공을 흉내 내지 않으며 가입을 안전하게 차단한다.",
    "자동화 테스트에서는 고정된 명시적 테스트 토큰만 허용하는 별도 테스트 어댑터를 사용한다.",
])

add_heading(doc, "5. 약관·동의 처리")
add_body(doc, "공급자 신규 가입자는 CUSTOMER와 PROVIDER 역할을 동시에 취득하므로 고객 필수약관, 공급자 필수약관 및 공통 필수약관의 활성 버전에 모두 동의해야 한다. 기존 고객이 공급자 역할을 추가할 때는 이미 동의한 고객 약관을 보존하고, 현재 유효한 공급자 필수약관 중 미동의 항목을 추가 동의받는다.")
add_bullets(doc, [
    "기존 공급자 계정 보정 마이그레이션은 고객 역할·프로필만 보완하며 고객 약관 동의를 임의 생성하지 않는다.",
    "고객 기능 최초 진입 시 현재 필수 고객 약관 미동의 여부를 확인하고 동의 화면으로 안내해야 한다.",
    "동의 이력은 버전, 시각, 출처, IP·User-Agent 등 기존 감사 의미를 유지한다.",
])

add_heading(doc, "6. 데이터 마이그레이션 및 안전성")
add_body(doc, "기존 공급자 전용 계정은 멱등 SQL 마이그레이션으로 보정한다. 활성 PROVIDER 역할을 가진 활성 사용자 중 CUSTOMER 역할이 없는 사용자에게만 역할을 추가하고, 공급자 프로필은 있으나 고객 프로필이 없는 사용자에게만 고객 프로필을 생성한다.")
add_bullets(doc, [
    "기존 사용자·공급자 프로필·지갑·원장·거래·신뢰도·개인정보·탈퇴 의미를 변경하지 않는다.",
    "고객 프로필 표시명은 대표자명 → 담당자명 → 상호 → 로그인 아이디 순으로 선택한다.",
    "롤백 시 고객 업무 데이터의 소유권 손상을 방지하기 위해 자동 삭제하지 않는다.",
    "운영 DB에 reset/delete를 수행하지 않고 배포 전 백업 및 사후 건수 검증을 실시한다.",
])

add_heading(doc, "7. 보안 및 권한 기준")
add_table(doc,
          ["통제", "기준"],
          [
              ["인증 쿠키", "HttpOnly·Secure·SameSite 및 허용 도메인 정책을 유지한다."],
              ["역할 권한", "CUSTOMER API와 PROVIDER API는 서버 권한검사를 각각 유지한다."],
              ["승인 분리", "PROVIDER 승인 전 공급자 업무는 차단하되 고객 기능은 허용한다."],
              ["세션 갱신", "역할 추가 성공 후 서버에서 역할을 재조회해 인증 쿠키를 재발급한다."],
              ["감사", "역할 부여·취소, 프로필 생성, 승인 변경, 인증 결과를 추적 가능하게 기록한다."],
          ], [4.1, 13.4])

add_heading(doc, "8. 구현 및 검증 체크리스트")
add_table(doc,
          ["검증 항목", "기대 결과"],
          [
              ["공급자 신규 가입", "CUSTOMER·PROVIDER 역할, 두 프로필, 지갑이 한 번만 생성됨"],
              ["고객의 공급자 등록", "기존 계정·고객 데이터 유지, PROVIDER만 추가, 세션 즉시 갱신"],
              ["기존 공급자 보정", "누락 계정만 CUSTOMER·고객 프로필 생성, 중복 0건"],
              ["본인인증 미연동", "고객·공급자 가입 모두 VERIFIED로 저장되지 않고 안전하게 차단"],
              ["사이트 전환", "공급자 홈·고객 홈 이동 후 재로그인 없이 권한별 화면 접근"],
              ["승인 전 공급자", "공급자 업무는 제한, 고객 기능은 정상"],
              ["홈 배너", "로그인 상태는 서비스 선택 안내, 비로그인은 로그인 안내"],
              ["회귀 검증", "Wallet/Fee/Trust/Privacy/Chat/Care/Interior/Emergency/탈퇴/차단 의미 유지"],
          ], [5.0, 12.5])

add_heading(doc, "9. 배포 순서")
add_bullets(doc, [
    "1) 운영 DB 백업과 현재 역할·프로필 중복 건수 확인",
    "2) API·프론트 정적 파일 배포 및 앱풀 재시작",
    "3) 통합계정 보정 마이그레이션 적용(멱등 실행)",
    "4) API health/readiness, CORS, 로그인, 역할 전환, 가입 차단 smoke test",
    "5) 외부 본인인증 업체 승인 후 운영 어댑터·Secret을 IIS 환경설정으로 연결",
    "6) 실제 인증 성공·실패·만료·번호불일치 시나리오를 검증한 뒤 가입 기능 공개",
])

add_callout(doc, "운영 전 남은 항목", "휴대폰 본인인증 업체 계약·승인, 운영 Secret 설정, 인증 콜백 검증, 고객 필수약관 최초 진입 Gate의 최종 확인이 필요하다. 이 항목이 완료되기 전에는 회원가입 성공을 가짜로 처리하지 않는다.", fill="FFF4E5", title_color=RED)

# Footer
for sec in doc.sections:
    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("수달 라이프 | 고객·공급자 통합계정 및 휴대폰 본인인증 정책 v1.0 | 내부 운영·개발 기준")
    set_run_font(r, size=7.5, color=GRAY)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
