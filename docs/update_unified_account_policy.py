from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt


ROOT = Path(r"D:\ALL_project\Soodal_Life_Project")
SOURCE = ROOT / "docs" / "output" / "수달라이프_고객공급자_통합계정_및_휴대폰본인인증_정책_v1.0.docx"
OUTPUT = ROOT / "docs" / "output" / "수달라이프_고객공급자_통합계정_및_휴대폰본인인증_정책_v1.1.docx"


def set_run_font(run, size: float, bold: bool = False) -> None:
    run.font.name = "맑은 고딕"
    run.font.size = Pt(size)
    run.bold = bold
    run._element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "맑은 고딕")


document = Document(SOURCE)

for section in document.sections:
    for paragraph in section.footer.paragraphs:
        for run in paragraph.runs:
            run.text = run.text.replace("정책 v1.0", "정책 v1.1")

for table in document.tables:
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                if "버전 1.0" in paragraph.text:
                    for run in paragraph.runs:
                        run.text = run.text.replace("버전 1.0", "버전 1.1")

page_break = document.add_paragraph()
page_break.add_run().add_break(WD_BREAK.PAGE)

heading = document.add_paragraph(style="Heading 1")
heading.paragraph_format.space_before = Pt(0)
heading.paragraph_format.space_after = Pt(10)
set_run_font(heading.add_run("10. 비밀번호 정책"), 16, True)

lead = document.add_paragraph()
lead.paragraph_format.space_after = Pt(8)
set_run_font(
    lead.add_run(
        "고객·공급자 통합계정의 비밀번호 규칙은 회원가입, 비밀번호 변경 및 비밀번호 재설정에 동일하게 적용한다."
    ),
    10.5,
)

items = [
    "비밀번호는 6자 이상이어야 한다.",
    "영문자, 숫자, 특수문자를 각각 1자 이상 포함해야 한다.",
    "영문 대문자는 필수가 아니며 소문자만 포함해도 다른 조건을 충족하면 사용할 수 있다.",
    "공백은 사용할 수 없으며, 비밀번호 확인값이 일치하지 않거나 규칙을 충족하지 않으면 요청을 거부한다.",
    "기존 비밀번호 해시는 임의 변경하거나 일괄 초기화하지 않는다. 새 비밀번호를 등록하는 시점부터 본 규칙을 적용한다.",
]

for text in items:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    set_run_font(paragraph.add_run(text), 10.5)

note = document.add_paragraph()
note.paragraph_format.space_before = Pt(8)
note.paragraph_format.space_after = Pt(0)
set_run_font(note.add_run("운영 유의사항  "), 10.5, True)
set_run_font(
    note.add_run(
        "휴대폰 본인인증 외부 연동이 완료되기 전에는 인증 성공을 임의로 생성하지 않고 NOT_INTEGRATED 및 Fail Closed 상태를 유지한다."
    ),
    10.5,
)

document.save(OUTPUT)
print(OUTPUT)
